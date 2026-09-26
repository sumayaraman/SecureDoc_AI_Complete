import re
from datetime import datetime
from pathlib import Path

def extract_excel_text(path: str, is_csv: bool = False) -> str:
    if is_csv:
        try:
            import csv
            lines = []
            with open(path, 'r', encoding='utf-8', errors='ignore') as f:
                reader = csv.reader(f)
                for row in reader:
                    row_str = ' | '.join(str(cell).strip() for cell in row if str(cell).strip())
                    if row_str:
                        lines.append(row_str)
            return '\n'.join(lines)
        except Exception:
            return ''
    
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        lines = []
        for sheetname in wb.sheetnames:
            ws = wb[sheetname]
            lines.append(f"Sheet: {sheetname}")
            for row in ws.iter_rows(values_only=True):
                non_empty = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if non_empty:
                    lines.append(' | '.join(non_empty))
        return '\n'.join(lines)
    except Exception:
        pass

    try:
        import zipfile
        import xml.etree.ElementTree as ET
        with zipfile.ZipFile(path) as z:
            shared_strings = []
            if 'xl/sharedStrings.xml' in z.namelist():
                tree = ET.fromstring(z.read('xl/sharedStrings.xml'))
                for t in tree.iter('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}t'):
                    shared_strings.append(t.text or '')
            
            lines = []
            for name in z.namelist():
                if name.startswith('xl/worksheets/sheet') and name.endswith('.xml'):
                    tree = ET.fromstring(z.read(name))
                    for row in tree.iter('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}row'):
                        cells = []
                        for c in row.iter('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}c'):
                            t = c.get('t')
                            val_elem = c.find('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}v')
                            if val_elem is not None and val_elem.text:
                                if t == 's' and int(val_elem.text) < len(shared_strings):
                                    cells.append(shared_strings[int(val_elem.text)])
                                else:
                                    cells.append(val_elem.text)
                        if cells:
                            lines.append(' | '.join(cells))
            return '\n'.join(lines)
    except Exception:
        return ''

def extract_text(path:str,mime:str,filename:str='')->str:
    ext = Path(filename or path).suffix.lower()
    if mime == 'application/pdf' or ext == '.pdf':
        try:
            import fitz
            doc=fitz.open(path); text='\n'.join(page.get_text() for page in doc)
            if text.strip(): return text
            # OCR scanned PDF pages when pytesseract + poppler-like rendering is available.
            return ocr_image_document(path, is_pdf=True)
        except Exception: return ''
    if mime in ('text/csv', 'application/csv') or ext == '.csv':
        return extract_excel_text(path, is_csv=True)
    if 'spreadsheet' in mime or 'excel' in mime or ext in ('.xlsx', '.xls'):
        return extract_excel_text(path, is_csv=False)
    if mime.startswith('image/') or ext in ('.jpg', '.jpeg', '.png'):
        return ocr_image_document(path, is_pdf=False)
    return ''

def ocr_image_document(path:str,is_pdf:bool=False)->str:
    try:
        import pytesseract
        from PIL import Image
        if is_pdf:
            import fitz
            doc=fitz.open(path); chunks=[]
            for page in doc:
                pix=page.get_pixmap(matrix=fitz.Matrix(2,2),alpha=False)
                img=Image.frombytes('RGB',[pix.width,pix.height],pix.samples)
                chunks.append(pytesseract.image_to_string(img))
            return '\n'.join(chunks)
        return pytesseract.image_to_string(Image.open(path))
    except Exception:
        return ''

def money(text:str, labels:list[str]):
    for label in labels:
        m=re.search(rf'\b{re.escape(label)}\b\s*[:#|–-]?\s*(?:৳|BDT|USD|\$)?\s*([0-9][0-9,]*(?:\.\d+)?)',text,re.I)
        if m:return float(m.group(1).replace(',',''))
    return 0.0

def classify_document(text: str, filename: str) -> dict:
    combined = f'{filename}\n{text}'.lower()
    cv_signals = ['curriculum vitae', 'resume', 'work experience', 'education', 'skills', 
                  'employment history', 'personal profile', 'academic background', 'publications', 'references available']
    cv_matches = sum(1 for s in cv_signals if s in combined)
    invoice_signals = [
        'invoice', 'bill to', 'billed to', 'tax invoice', 'invoice no', 'invoice number', 
        'invoice #', 'due date', 'subtotal', 'total due', 'amount due', 'payment terms',
        'remit to', 'purchase order', 'po #', 'challan', 'vat reg', 'tin'
    ]
    receipt_signals = [
        'receipt', 'sales receipt', 'cash receipt', 'pos receipt', 'merchant',
        'cashier', 'tender', 'change due', 'transaction no'
    ]
    invoice_matches = sum(1 for s in invoice_signals if s in combined)
    receipt_matches = sum(1 for s in receipt_signals if s in combined)
    has_money = bool(re.search(r'(?:৳|bdt|usd|\$|eur|€|total)\s*[:#-]?\s*[0-9]+(?:\.[0-9]+)?', combined, re.I))

    if cv_matches >= 2 and invoice_matches == 0 and receipt_matches == 0:
        return {'is_supported': False, 'doc_type': 'resume', 'reason': 'Detected as a Resume / CV, not an invoice or receipt.'}
    if invoice_matches >= 1:
        return {'is_supported': True, 'doc_type': 'invoice', 'confidence': min(95.0, 50.0 + invoice_matches * 15.0)}
    if receipt_matches >= 1:
        return {'is_supported': True, 'doc_type': 'receipt', 'confidence': min(90.0, 50.0 + receipt_matches * 15.0)}
    if has_money and any(k in combined for k in ['vendor', 'supplier', 'date', 'total', 'item', 'price', 'qty', 'tax']):
        return {'is_supported': True, 'doc_type': 'invoice', 'confidence': 60.0}
    if any(k in filename.lower() for k in ['invoice', 'bill', 'receipt', 'inv_', 'rec_', 'sheet', 'billing', 'expense', 'order', 'sales']):
        return {'is_supported': True, 'doc_type': 'invoice', 'confidence': 50.0}
    return {'is_supported': False, 'doc_type': 'unsupported', 'reason': "We couldn't identify this file as a supported invoice or receipt."}

def extract_simple_items(lines: list[str]) -> list[dict]:
    items = []
    pattern_space = re.compile(
        r'^\s*([A-Za-z0-9\s.,&/\(\)-]{2,60}?)\s+(\d+(?:\.\d+)?)\s+(?:৳|BDT|USD|\$)?\s*(\d+(?:,\d+)*(?:\.\d+)?)\s+(?:৳|BDT|USD|\$)?\s*(\d+(?:,\d+)*(?:\.\d+)?)\s*$',
        re.I
    )
    pattern_pipe = re.compile(
        r'^\s*([A-Za-z0-9\s.,&/\(\)-]{2,60}?)\s*\|\s*(\d+(?:\.\d+)?)\s*\|\s*(?:৳|BDT|USD|\$)?\s*(\d+(?:,\d+)*(?:\.\d+)?)\s*\|\s*(?:৳|BDT|USD|\$)?\s*(\d+(?:,\d+)*(?:\.\d+)?)\s*$',
        re.I
    )
    for line in lines:
        m = pattern_pipe.match(line) or pattern_space.match(line)
        if m:
            desc = m.group(1).strip()
            if any(h in desc.lower() for h in ['description', 'unit price', 'subtotal', 'grand total', 'tax']):
                continue
            qty = float(m.group(2))
            price = float(m.group(3).replace(',', ''))
            tot = float(m.group(4).replace(',', ''))
            if qty > 0 and (price > 0 or tot > 0):
                items.append({'description': desc, 'quantity': qty, 'unit_price': price, 'total': tot})
    return items

def parse_invoice(text:str,filename:str)->dict:
    clean=' '.join(text.split()); lines=[x.strip() for x in text.splitlines() if x.strip()]
    vendor=''
    for line in lines:
        m=re.search(r'^(?:vendor|supplier|seller|from)\s*[:#|–-]\s*([^|#\n]+)$',line,re.I)
        if m:
            vendor=m.group(1).strip()
            break
    if not vendor:
        m=re.search(r'(?:vendor|supplier|seller|from)\s*[:#|–-]\s*([A-Za-z0-9 &.,_-]{2,80})',clean,re.I)
        if m: vendor=m.group(1).strip()
    if not vendor:
        for line in lines[:5]:
            if not any(h in line.lower() for h in ['invoice', 'sheet', 'date', 'total', 'bill to', 'tax']):
                vendor=line.split('|')[0].strip()[:120]
                break
    if not vendor: vendor=lines[0][:120] if lines else 'Unknown Vendor'

    num=''
    m=re.search(r'invoice\s*(?:number|no\.?|#)\s*[:#|–-]?\s*([A-Za-z0-9/_-]{2,})',clean,re.I)
    if m:
        num=m.group(1).strip()
    else:
        m=re.search(r'invoice\s*[:#|–-]\s*([A-Za-z0-9/_-]{2,})',clean,re.I)
        if m: num=m.group(1).strip()

    date=''
    m=re.search(r'(?:invoice\s*date|date)\s*[:#|–-]?\s*(\d{1,4}[-/.]\d{1,2}[-/.]\d{1,4})',clean,re.I)
    if m:date=m.group(1)

    due_date=''
    m=re.search(r'(?:due\s*date)\s*[:#|–-]?\s*(\d{1,4}[-/.]\d{1,2}[-/.]\d{1,4})',clean,re.I)
    if m:due_date=m.group(1)

    currency='BDT' if re.search(r'৳|\bBDT\b',clean,re.I) else 'USD' if re.search(r'\$|\bUSD\b',clean,re.I) else 'BDT'
    subtotal=money(clean,['subtotal','sub total']); discount=money(clean,['discount']); tax=money(clean,['tax','vat']); total=money(clean,['grand total','total amount','total due','total'])
    items = extract_simple_items(lines)
    if items and subtotal == 0:
        subtotal = sum(i['total'] for i in items)
    if total==0 and subtotal: total=subtotal-discount+tax
    confidence = 85.0 if (vendor and total > 0 and num) else 70.0 if (total > 0) else 50.0
    return {'vendor_name':vendor,'invoice_number':num or filename.rsplit('.',1)[0],'invoice_date':date or datetime.utcnow().date().isoformat(),'due_date':due_date,'subtotal':subtotal,'discount':discount,'tax':tax,'total':total,'currency':currency,'items':items,'confidence':confidence}
